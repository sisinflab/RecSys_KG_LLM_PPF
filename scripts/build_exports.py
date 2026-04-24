from pathlib import Path
from docx import Document
from openpyxl import load_workbook
import csv

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / 'source_materials' / 'History_20_RecSys_Combined.xlsx'
OUT = ROOT / 'data'

wb = load_workbook(SRC, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
header = [str(h).strip() for h in rows[0]]
records = []
for row in rows[1:]:
    rec = {header[i]: row[i] for i in range(len(header))}
    doi = str(rec.get('doi') or '').strip()
    rec['pdf_file_hint'] = doi.split('10.1145/')[-1] + '.pdf' if doi.startswith('10.1145/') else ''
    records.append(rec)

with (OUT / 'papers_master.csv').open('w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f)
    writer.writerow(['year','title','venue','doi','pdf_file_hint','authors','found_by_keywords','dblp_url'])
    for r in records:
        writer.writerow([
            r.get('year'), r.get('title'), r.get('venue'), r.get('doi'), r.get('pdf_file_hint'),
            r.get('authors'), r.get('found_by_keywords'), r.get('dblp_url')
        ])

with (OUT / 'bibliography.bib').open('w', encoding='utf-8') as f:
    entries = [str(r.get('bibtex') or '').strip() for r in records if str(r.get('bibtex') or '').strip()]
    f.write('

'.join(entries) + '
')

print('Rebuilt CSV and BibTeX exports.')
